import asyncio
import json
import re
from pathlib import Path
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Emu

BUNDLE = Path("/Users/jeffryhernandez/Documents/kserve/talk/v5-bundle")
HTML = BUNDLE / "Deploying AI-Native Apps on Kubernetes.html"
RENDER_DIR = Path("/Users/jeffryhernandez/Documents/kserve/talk/_render")
PPTX_OUT = Path("/Users/jeffryhernandez/Documents/kserve/talk/CNCG-Cali-v5.pptx")

RENDER_DIR.mkdir(parents=True, exist_ok=True)


async def render_slides() -> int:
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        ctx = await browser.new_context(
            viewport={"width": 1920, "height": 1080},
            device_scale_factor=2,
        )
        page = await ctx.new_page()
        await page.goto(f"file://{HTML}")
        await page.wait_for_function(
            "typeof document.querySelector('deck-stage')?.goTo === 'function'"
        )
        await page.evaluate(
            "document.querySelector('deck-stage').setAttribute('noscale','')"
        )
        await page.evaluate("""
            const sr = document.querySelector('deck-stage').shadowRoot;
            const style = document.createElement('style');
            style.textContent = '.export-hidden { display: none !important; }';
            sr.appendChild(style);
        """)
        await page.wait_for_timeout(300)
        total = await page.evaluate(
            "document.querySelectorAll('deck-stage > section').length"
        )
        print(f"slides detected: {total}")
        for i in range(total):
            await page.evaluate(
                f"document.querySelector('deck-stage').goTo({i})"
            )
            await page.wait_for_function(
                f"document.querySelectorAll('deck-stage > section')[{i}]"
                ".hasAttribute('data-deck-active')"
            )
            await page.wait_for_timeout(200)
            out = RENDER_DIR / f"slide-{i:02d}.png"
            await page.screenshot(
                path=str(out),
                clip={"x": 0, "y": 0, "width": 1920, "height": 1080},
            )
            print(f"  rendered slide {i+1:02d}/{total} -> {out.name}")
        await browser.close()
    return total


def load_speaker_notes() -> list[str]:
    raw = HTML.read_text()
    block = re.search(
        r'<script type="application/json" id="speaker-notes">(.*?)</script>',
        raw, re.DOTALL,
    ).group(1)
    notes = json.loads(block)
    # The JSON ships 43 notes for 40 slides. Indices 10, 19, 37 are
    # duplicate "Sección N" openers that have no corresponding slide
    # (the slide before each one already gets a transition note).
    drop = {10, 19, 37}
    return [n for i, n in enumerate(notes) if i not in drop]


def build_pptx(total: int) -> None:
    notes = load_speaker_notes()
    if len(notes) != total:
        print(
            f"WARNING: {len(notes)} notes vs {total} slides — "
            "alignment may be off. Check the drop set."
        )
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333"
    prs.slide_height = Emu(6858000)   # 7.5"
    blank = prs.slide_layouts[6]
    for i in range(total):
        png = RENDER_DIR / f"slide-{i:02d}.png"
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(png), 0, 0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        if i < len(notes):
            slide.notes_slide.notes_text_frame.text = notes[i]
    prs.save(str(PPTX_OUT))
    print(f"\nwrote {PPTX_OUT}  ({len(notes)} speaker notes attached)")


async def main():
    total = await render_slides()
    build_pptx(total)


asyncio.run(main())
